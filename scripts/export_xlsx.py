"""Exports site/data/opportunities.json (the source of truth) to Funding_Radar_Registry.xlsx for reading and hand-checks."""
import json, sys
import os
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE); sys.path.insert(0, HERE); os.chdir(ROOT)
_d = json.load(open("site/data/opportunities.json", encoding="utf-8"))
REGISTRY = _d["opportunities"]; DISCOVERY_FEEDS = _d.get("discovery_feeds", []); EXCLUDED = _d.get("excluded", []); CHECKED = _d["meta"]["generated"]
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Typical application months (1-12) used for the annual calendar; [] = year-round/irregular
MONTHS = {
 "daad-research-stays":[9,10], "daad-grow":[3,4], "swiss-research-fellowship":[9,10,11,12], "csc-split-site":[10,11,12],
 "ictp-step":[1,2,3], "owsd-phd":[2,3,4,5], "embo-seg":[], "bif-travel":[], "cob-travelling":[], "imu-abel":[5,6],
 "erasmus-doctoral-short":[], "ieee-qce-travel":[5,6,7,8], "qcrypt-support":[6,7], "spie-student-support":[],
 "optica-milton-chang":[4,5], "ieee-photonics-travel":[], "google-conf-scholarships":[], "sigplan-pac":[],
 "acm-w-scholarship":[], "ais-bocconi":[], "aps-fip-ds":[11,12], "aps-eps-ictp":[2,3], "ibro-sfn":[5,6],
 "icm-travel":[], "csndsp-travel":[5], "esg-travel":[], "ecsr-travel":[2], "sssr-travel":[], "iamcr-bersay":[],
 "eshms-travel":[], "wenner-gren-dfg":[4,5,10,11], "eisa-fieldwork":[4,5], "codesria-small-grants":[],
 "ssrc-nextgen-africa":[12,1], "inlaks-research-travel":[7,8],
}
TYPE_LABEL = {"research_stay":"Research stay","mobility":"Mobility","conference_travel":"Conference travel",
              "fieldwork":"Fieldwork","thesis_completion":"Thesis completion","small_grant":"Small grant"}
STATUS_LABEL = {"open":"Open now","upcoming":"Opens soon (recurring)","rolling":"Rolling / year-round",
                "monitor":"Recurring – watch for next call","paused":"Paused","closed":"Closed this cycle"}

for r in REGISTRY:
    r.setdefault("months", MONTHS.get(r["id"], []))
    r["type_label"] = TYPE_LABEL[r["type"]]
    r["status_label"] = STATUS_LABEL[r["status"]]

meta = {"generated": CHECKED, "count": len(REGISTRY), "version": 1,
        "verification_counts": {v: sum(1 for r in REGISTRY if r["verification"] == v) for v in ["official","secondary","needs_check"]}}

# ---------------- Excel ----------------
wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3A5F")
wrap = Alignment(wrap_text=True, vertical="top")
def header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i); cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
    ws.freeze_panes = "B2"

ws = wb.active; ws.title = "Registry"
cols = ["ID","Programme","Provider","Type","Levels","Fields","Nationality rule","Enrolment rule","Host","Amount","Covers",
        "Paid how","Membership needed","Recurrence","Typical window","Next deadline","Status","Developing-country priority",
        "Official URL","URL checked","Verification","Last checked","Notes"]
header(ws, cols, [20,40,28,16,10,22,34,34,20,32,30,12,22,22,30,13,22,12,45,10,12,12,50])
for r in REGISTRY:
    ws.append([r["id"], r["name"], r["provider"], r["type_label"], ", ".join(r["levels"]), ", ".join(r["fields"]),
               f'{r["nat_rule"]["type"]}: {r["nat_rule"].get("note","")}'.strip(": "),
               f'{r["enrol_rule"]["type"]}: {r["enrol_rule"].get("note","")}'.strip(": "),
               r["host"], r["amount"], r["covers"], r["payment"], r["membership"], r["recurrence"], r["window"],
               r["next_deadline"] or "", r["status_label"], "Yes" if r["dev_priority"] else "No", r["url"],
               "Yes" if r["url_checked"] else "No", r["verification"], r["checked"], r["notes"]])
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = wrap
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(REGISTRY)+1}"
dv = DataValidation(type="list", formula1='"official,secondary,needs_check"', allow_blank=False); ws.add_data_validation(dv)
dv.add(f"U2:U{len(REGISTRY)+200}")

# Calendar sheet
cal = wb.create_sheet("Annual calendar")
mn = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
header(cal, ["Programme","Type"]+mn+["Note"], [42,16]+[5]*12+[40])
on = PatternFill("solid", fgColor="F2A541")
for r in sorted(REGISTRY, key=lambda x: (x["months"][0] if x["months"] else 99)):
    cal.append([r["name"], r["type_label"]] + ["●" if (i+1) in r["months"] else "" for i in range(12)] +
               ["Year-round / per event" if not r["months"] else r["window"]])
    rr = cal.max_row
    for i in range(12):
        if (i+1) in r["months"]: cal.cell(row=rr, column=3+i).fill = on
    for c in cal[rr]: c.alignment = wrap

fd = wb.create_sheet("Discovery feeds")
header(fd, ["Feed","URL","Kind","Rule"], [40,60,24,60])
for f in DISCOVERY_FEEDS:
    fd.append([f["name"], f["url"], f["kind"], "Use only to find programmes; always cite the provider's own page"])

ex = wb.create_sheet("Excluded (out of scope)")
header(ex, ["Programme","Reason excluded","URL"], [50,60,50])
for e in EXCLUDED: ex.append([e["name"], e["reason"], e["url"]])

qa = wb.create_sheet("QA log")
header(qa, ["Week starting","Items reviewed","In scope (%)","Dead links (%)","Deadline correct (%)","Eligibility correct (%)",
            "Errors reported by users","Fixed within 48h (%)","Reviewer","Notes"], [14,12,12,12,14,14,14,14,14,40])
qa.append(["2026-11-02","","","","","","","","","First weekly hand-check: sample 20 items"])
qa.append(["Targets →","≥20","≥95","≤2","≥98","≥95","","100","",""])

dd = wb.create_sheet("Data dictionary")
header(dd, ["Field","Meaning","Allowed values"], [22,70,50])
for row in [
 ("type","Kind of support","research_stay, mobility, conference_travel, fieldwork, thesis_completion, small_grant"),
 ("levels","Study level eligible","MSc, PhD"),
 ("nat_rule","Nationality condition","any | developing (provider's list) | list | region"),
 ("enrol_rule","Where the student must be enrolled","any | developing | list | region"),
 ("payment","How money arrives (key pain point)","upfront | reimbursed | stipend | varies"),
 ("status","Current cycle state","open, upcoming, rolling, monitor, paused, closed"),
 ("verification","Strength of confirmation","official (provider page), secondary, needs_check"),
 ("url_checked","Has the URL itself been opened and confirmed","Yes / No"),
 ("dev_priority","Programme targets or prioritises developing-country applicants","Yes / No"),
 ("months","Typical application months, for the annual calendar","1–12, empty = year-round"),
]: dd.append(list(row))
wb.save("Funding_Radar_Registry.xlsx")
print("built", meta)
